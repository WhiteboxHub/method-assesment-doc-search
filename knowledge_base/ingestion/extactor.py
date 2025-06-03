
from pptx import Presentation
import os
import json

def extract_pptx_info(filepath):
    prs = Presentation(filepath)

    # Metadata
    core_props = prs.core_properties
    metadata = {
        "filename": os.path.basename(filepath),
        "title": core_props.title,
        "author": core_props.author
    }

    # Extract text slide-by-slide
    slides = []
    for i, slide in enumerate(prs.slides):
        text_list = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                text_list.append(shape.text.strip())

        slide_data = {
            "slide_number": i + 1,
            "text": "\n".join(text_list)
        }
        slides.append(slide_data)

    return {
        "metadata": metadata,
        "slides": slides
    }

def save_to_json(data, output_path):
    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=4, ensure_ascii=False)

def process_all_pptx_in_folder(folder_path, output_folder):
    os.makedirs(output_folder, exist_ok=True)

    for file in os.listdir(folder_path):
        if file.endswith(".pptx"):
            file_path = os.path.join(folder_path, file)
            print(f"Processing: {file_path}")
            extracted_data = extract_pptx_info(file_path)

            output_filename = os.path.splitext(file)[0] + ".json"
            output_path = os.path.join(output_folder, output_filename)

            save_to_json(extracted_data, output_path)
            print(f"Saved to: {output_path}")

# Example usage
if __name__ == "__main__":
    input_folder = r"knowledge_base\corpus\raw"        # <-- Update if needed
    output_folder = r"knowledge_base\corpus\JsonOut"   # <-- Output JSONs will be saved here

    process_all_pptx_in_folder(input_folder, output_folder)