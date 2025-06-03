from pptx import Presentation
from unstructured.partition.text import partition_text

def extract_unknow_slide(pptx_path):
    print(pptx_path)
    prs = Presentation(pptx_path)

    slides_data = []

    for i,slide in enumerate(prs.slides):
        raw_texts = []

        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.has_text_frame:
                text = shape.text.strip()
                if text:
                    raw_texts.append(text)
        
        full_text = "\n".join(raw_texts)

        elemetns = partition_text(text=full_text)

        slides_data.append({
            "slideno":i+1,
            "raw_text": full_text,
            "elements":[{"category":el.category,"text":el.text} for el in elemetns]
        })
    return slides_data